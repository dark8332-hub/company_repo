"""반입 런북을 마크다운 원본에서 Word(.docx)와 HTML 로 만든다.

    .venv/bin/python tools/build_runbook.py            # docs/폐쇄망-반입-런북.{docx,html}
    .venv/bin/python tools/build_runbook.py <원본.md>   # 다른 문서

마크다운이 원본이고 나머지는 산출물이다. 손으로 고치지 않는다 — 산출물만 고치면 다음 생성에서
사라지고, 그 전까지는 세 파일이 서로 다른 말을 한다(1.1.4 시절 .docx 가 그랬다).

지원하는 구문은 런북이 실제로 쓰는 것뿐이다: 제목(#~####), 문단(줄바꿈으로 이어진 것 포함),
표, 코드 블록, 인용(> 로 시작, 여러 줄), 순서·비순서 목록(들여쓴 이어지는 줄 포함),
이미지(![캡션](경로)), 구분선, 인라인 `코드`·**굵게**·*기울임*.

HTML 은 이미지를 data URI 로 안에 넣고 외부 글꼴·스크립트를 참조하지 않는다. 폐쇄망에서 열리는
문서가 밖으로 무언가를 가지러 나가면 실패가 아니라 지연으로 나타나 원인을 찾기 어렵다.
"""
import base64
import html
import mimetypes
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DOCS = Path(__file__).resolve().parent.parent / "docs"
DEFAULT_SOURCE = DOCS / "폐쇄망-반입-런북.md"


# --- 마크다운 → 블록 목록 -----------------------------------------------------------------------

TABLE_RULE = re.compile(r"^\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?$")
IMAGE = re.compile(r"^!\[(?P<alt>[^\]]*)\]\((?P<src>[^)]+)\)$")
HEADING = re.compile(r"^(#{1,4})\s+(.*)$")
ORDERED = re.compile(r"^(\d+)\.\s+(.*)$")
BULLET = re.compile(r"^[-*]\s+(.*)$")


def split_row(line):
    cells, current, in_code = [], [], False
    for ch in line.strip().strip("|"):
        if ch == "`":
            in_code = not in_code
        if ch == "|" and not in_code:
            cells.append("".join(current).strip())
            current = []
        else:
            current.append(ch)
    cells.append("".join(current).strip())
    return cells


def parse(text):
    """블록 목록: (종류, 내용). 문단·목록 항목의 이어지는 줄은 한 덩어리로 합친다."""
    lines = text.splitlines()
    blocks, i = [], 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped == "---":
            blocks.append(("rule", None))
            i += 1
            continue
        if stripped.startswith("```"):
            code, i = [], i + 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            blocks.append(("code", "\n".join(code)))
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and TABLE_RULE.match(lines[i + 1].strip()):
            header = split_row(stripped)
            rows, i = [], i + 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            blocks.append(("table", (header, rows)))
            continue
        m = HEADING.match(stripped)
        if m:
            blocks.append(("heading", (len(m.group(1)), m.group(2).strip())))
            i += 1
            continue
        m = IMAGE.match(stripped)
        if m:
            blocks.append(("image", (m.group("alt"), m.group("src"))))
            i += 1
            continue
        if stripped.startswith(">"):
            quote = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip()[1:].strip())
                i += 1
            blocks.append(("quote", " ".join(q for q in quote if q)))
            continue
        if BULLET.match(stripped) or ORDERED.match(stripped):
            items, ordered = [], bool(ORDERED.match(stripped))
            while i < len(lines):
                cur = lines[i]
                cs = cur.strip()
                m = ORDERED.match(cs) if ordered else BULLET.match(cs)
                if m:
                    items.append(m.group(2) if ordered else m.group(1))
                    i += 1
                elif cs and cur.startswith("  ") and items:
                    items[-1] += " " + cs
                    i += 1
                else:
                    break
            blocks.append(("list", (ordered, items)))
            continue
        para = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i]
            ns = nxt.strip()
            if (not ns or ns == "---" or ns.startswith(("```", "|", ">", "#", "![")) or BULLET.match(ns) or ORDERED.match(ns)):
                break
            para.append(ns)
            i += 1
        blocks.append(("para", " ".join(para)))
    return blocks


INLINE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*\s][^*]*\*)")


def inline_runs(text):
    """(종류, 글) 목록. 종류는 text / code / bold / italic."""
    out = []
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("`"):
            out.append(("code", piece[1:-1]))
        elif piece.startswith("**"):
            out.append(("bold", piece[2:-2]))
        elif piece.startswith("*"):
            out.append(("italic", piece[1:-1]))
        else:
            out.append(("text", piece))
    return out


def plain(text):
    return "".join(t for _, t in inline_runs(text))


# --- Word --------------------------------------------------------------------------------------

BODY_FONT = "맑은 고딕"
MONO_FONT = "Consolas"
INK = RGBColor(0x15, 0x1B, 0x24)
SOFT = RGBColor(0x4A, 0x56, 0x65)
NAVY = RGBColor(0x1E, 0x3A, 0x5F)


def shade(element, color):
    tag = OxmlElement("w:shd")
    tag.set(qn("w:val"), "clear")
    tag.set(qn("w:fill"), color)
    element.append(tag)


def set_font(run, font):
    run.font.name = font
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font)


def write_inline(paragraph, text, size=10.5, color=INK):
    for kind, piece in inline_runs(text):
        run = paragraph.add_run(piece)
        run.font.size = Pt(size)
        run.font.color.rgb = color
        if kind == "code":
            set_font(run, MONO_FONT)
            run.font.size = Pt(size - 0.5)
            run.font.color.rgb = NAVY
        else:
            set_font(run, BODY_FONT)
            run.bold = kind == "bold"
            run.italic = kind == "italic"


def build_docx(blocks, source_dir, out_path):
    document = Document()
    normal = document.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.font.name = BODY_FONT
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), BODY_FONT)
    for section in document.sections:
        section.left_margin = section.right_margin = Cm(2.2)
        section.top_margin = section.bottom_margin = Cm(2.0)
    text_width = document.sections[0].page_width - document.sections[0].left_margin - document.sections[0].right_margin
    figure = 0

    for kind, body in blocks:
        if kind == "rule":
            continue
        if kind == "heading":
            level, title = body
            paragraph = document.add_heading(level=level)
            paragraph.paragraph_format.space_before = Pt(18 if level <= 2 else 10)
            paragraph.paragraph_format.space_after = Pt(6)
            run = paragraph.add_run(plain(title))
            set_font(run, BODY_FONT)
            run.font.color.rgb = NAVY
            run.font.size = Pt({1: 20, 2: 15, 3: 12.5, 4: 11}[level])
            if level == 1:
                # 표지 역할: 제목 뒤에 여백
                paragraph.paragraph_format.space_after = Pt(10)
            continue
        if kind == "para":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(7)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            write_inline(paragraph, body)
            continue
        if kind == "code":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.4)
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(8)
            shade(paragraph._p.get_or_add_pPr(), "F2F4F7")
            run = paragraph.add_run(body)
            set_font(run, MONO_FONT)
            run.font.size = Pt(9)
            continue
        if kind == "quote":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.5)
            paragraph.paragraph_format.space_after = Pt(8)
            shade(paragraph._p.get_or_add_pPr(), "FBF3E6")
            write_inline(paragraph, body, 10, SOFT)
            continue
        if kind == "list":
            ordered, items = body
            for item in items:
                paragraph = document.add_paragraph(style="List Number" if ordered else "List Bullet")
                paragraph.paragraph_format.space_after = Pt(3)
                write_inline(paragraph, item)
            continue
        if kind == "table":
            header, rows = body
            headless = not any(header)
            table = document.add_table(rows=0 if headless else 1, cols=len(header))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            if not headless:
                for cell, text in zip(table.rows[0].cells, header):
                    shade(cell._tc.get_or_add_tcPr(), "E8ECF1")
                    cell.paragraphs[0].text = ""
                    write_inline(cell.paragraphs[0], text or " ", 9.5)
                    for run in cell.paragraphs[0].runs:
                        run.bold = True
            for row in rows:
                cells = table.add_row().cells
                for cell, text in zip(cells, row + [""] * (len(header) - len(row))):
                    cell.paragraphs[0].text = ""
                    write_inline(cell.paragraphs[0], text or " ", 9.5)
            document.add_paragraph().paragraph_format.space_after = Pt(2)
            continue
        if kind == "image":
            alt, src = body
            path = (source_dir / src).resolve()
            if not path.exists():
                raise FileNotFoundError(f"그림 파일이 없습니다: {src}")
            figure += 1
            document.add_picture(str(path), width=text_width)
            document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption = document.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.paragraph_format.space_after = Pt(10)
            write_inline(caption, f"그림 {figure}. {alt}", 9, SOFT)
            continue
        raise ValueError(kind)

    document.save(str(out_path))


# --- HTML --------------------------------------------------------------------------------------

CSS = """
:root { color-scheme: light dark;
  --paper:#f4f6f8; --surface:#fff; --ink:#151b24; --soft:#4a5665; --rule:#d7dde5; --navy:#1e3a5f;
  --note:#fbf3e6; --note-rule:#e3c79a; --code-bg:#141b26; --code-ink:#e2e8f2; --inline-bg:#eef1f5; --head:#e8ecf1; }
@media (prefers-color-scheme: dark) { :root:not([data-theme="light"]) {
  --paper:#0f141b; --surface:#161d27; --ink:#e3e9f1; --soft:#a9b5c4; --rule:#2b3542; --navy:#8fb4e0;
  --note:#2a2418; --note-rule:#6b5530; --code-bg:#0b1017; --code-ink:#dde5f0; --inline-bg:#1c242f; --head:#1c2531; } }
* { box-sizing:border-box }
body { margin:0; background:var(--paper); color:var(--ink);
  font:15px/1.7 -apple-system,"Segoe UI","Malgun Gothic","Apple SD Gothic Neo","Noto Sans KR",sans-serif; }
.page { max-width:920px; margin:0 auto; padding:40px 24px 80px; }
article { background:var(--surface); border:1px solid var(--rule); border-radius:12px; padding:40px 48px; }
h1 { font-size:2rem; margin:0 0 .3em; color:var(--navy); line-height:1.25 }
h2 { font-size:1.45rem; margin:2.2em 0 .6em; padding-top:1.2em; border-top:1px solid var(--rule); color:var(--navy) }
h3 { font-size:1.12rem; margin:1.8em 0 .5em }
h4 { font-size:1rem; margin:1.4em 0 .4em }
p { margin:.6em 0 }
code { font-family:ui-monospace,"Cascadia Mono",Consolas,"D2Coding",monospace; font-size:.9em;
  background:var(--inline-bg); padding:.1em .35em; border-radius:4px }
pre { background:var(--code-bg); color:var(--code-ink); padding:14px 18px; border-radius:8px; overflow-x:auto;
  font-size:.86rem; line-height:1.55 }
pre code { background:none; padding:0; color:inherit; font-size:inherit }
blockquote { margin:1em 0; padding:10px 16px; background:var(--note); border-left:4px solid var(--note-rule);
  border-radius:0 8px 8px 0; color:var(--soft) }
blockquote p { margin:0 }
table { border-collapse:collapse; width:100%; margin:.8em 0 1.2em; font-size:.92rem }
.table-wrap { overflow-x:auto }
th, td { border:1px solid var(--rule); padding:7px 10px; vertical-align:top; text-align:left }
th { background:var(--head); font-weight:600 }
figure { margin:1.2em 0 1.6em; text-align:center }
figure img { max-width:100%; border:1px solid var(--rule); border-radius:8px }
figcaption { font-size:.86rem; color:var(--soft); margin-top:.5em }
ul, ol { padding-left:1.5em } li { margin:.25em 0 }
hr { border:0; height:0; margin:0 }
nav.toc { margin:1.2em 0 2em; padding:14px 20px; background:var(--inline-bg); border-radius:8px; font-size:.93rem }
nav.toc ol { margin:.3em 0 0; padding-left:1.3em; columns:2; column-gap:2em }
nav.toc a { color:var(--navy); text-decoration:none } nav.toc a:hover { text-decoration:underline }
.meta { color:var(--soft); font-size:.9rem }
@media (max-width:600px) { article { padding:24px 18px } nav.toc ol { columns:1 } }
@media print { body { background:#fff } article { border:0; padding:0 } h2 { break-before:page } h2:first-of-type { break-before:auto } figure { break-inside:avoid } }
"""


def inline_html(text):
    out = []
    for kind, piece in inline_runs(text):
        esc = html.escape(piece)
        if kind == "code":
            out.append(f"<code>{esc}</code>")
        elif kind == "bold":
            out.append(f"<strong>{esc}</strong>")
        elif kind == "italic":
            out.append(f"<em>{esc}</em>")
        else:
            out.append(esc)
    return "".join(out)


def slug(text, used):
    base = re.sub(r"[^\w가-힣]+", "-", plain(text)).strip("-").lower() or "s"
    candidate, n = base, 2
    while candidate in used:
        candidate, n = f"{base}-{n}", n + 1
    used.add(candidate)
    return candidate


def build_html(blocks, source_dir, out_path):
    title = next((plain(b[1]) for k, b in blocks if k == "heading" and b[0] == 1), out_path.stem)
    used, toc, parts, figure = set(), [], [], 0
    for kind, body in blocks:
        if kind == "rule":
            parts.append("<hr>")
        elif kind == "heading":
            level, text = body
            anchor = slug(text, used)
            if level == 2:
                toc.append((anchor, plain(text)))
            parts.append(f'<h{level} id="{anchor}">{inline_html(text)}</h{level}>')
        elif kind == "para":
            parts.append(f"<p>{inline_html(body)}</p>")
        elif kind == "code":
            parts.append(f"<pre><code>{html.escape(body)}</code></pre>")
        elif kind == "quote":
            parts.append(f"<blockquote><p>{inline_html(body)}</p></blockquote>")
        elif kind == "list":
            ordered, items = body
            tag = "ol" if ordered else "ul"
            parts.append(f"<{tag}>" + "".join(f"<li>{inline_html(i)}</li>" for i in items) + f"</{tag}>")
        elif kind == "table":
            header, rows = body
            head = "" if not any(header) else "<thead><tr>" + "".join(f"<th>{inline_html(c)}</th>" for c in header) + "</tr></thead>"
            body_rows = "".join("<tr>" + "".join(f"<td>{inline_html(c)}</td>" for c in r + [""] * (len(header) - len(r))) + "</tr>" for r in rows)
            parts.append(f'<div class="table-wrap"><table>{head}<tbody>{body_rows}</tbody></table></div>')
        elif kind == "image":
            alt, src = body
            path = (source_dir / src).resolve()
            if not path.exists():
                raise FileNotFoundError(f"그림 파일이 없습니다: {src}")
            figure += 1
            mime = mimetypes.guess_type(str(path))[0] or "image/png"
            data = base64.b64encode(path.read_bytes()).decode("ascii")
            parts.append(f'<figure><img src="data:{mime};base64,{data}" alt="{html.escape(alt)}" loading="lazy">'
                         f"<figcaption>그림 {figure}. {inline_html(alt)}</figcaption></figure>")
    # 목차는 첫 h1 바로 뒤에 둔다
    toc_html = ""
    if toc:
        toc_html = '<nav class="toc"><strong>차례</strong><ol>' + "".join(f'<li><a href="#{a}">{html.escape(t)}</a></li>' for a, t in toc) + "</ol></nav>"
    h1_index = next((n for n, p in enumerate(parts) if p.startswith("<h1")), None)
    if h1_index is not None:
        # h1 뒤의 요약표(있으면)까지 지나서 목차를 넣는다
        insert_at = h1_index + 1
        while insert_at < len(parts) and not parts[insert_at].startswith("<h2"):
            insert_at += 1
        parts.insert(insert_at, toc_html)
    document = (f"<!doctype html>\n<html lang=\"ko\">\n<head>\n<meta charset=\"utf-8\">\n"
                f"<meta name=\"viewport\" content=\"width=device-width, initial-scale=1\">\n"
                f"<title>{html.escape(title)}</title>\n<style>{CSS}</style>\n</head>\n<body>\n"
                f"<div class=\"page\"><article>\n{chr(10).join(parts)}\n</article></div>\n</body>\n</html>\n")
    out_path.write_text(document, encoding="utf-8")


def main():
    source = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else DEFAULT_SOURCE
    blocks = parse(source.read_text(encoding="utf-8"))
    docx_path = source.with_suffix(".docx")
    html_path = source.with_suffix(".html")
    build_docx(blocks, source.parent, docx_path)
    build_html(blocks, source.parent, html_path)
    images = sum(1 for k, _ in blocks if k == "image")
    print(f"{docx_path.name}  {docx_path.stat().st_size // 1024}KB")
    print(f"{html_path.name}  {html_path.stat().st_size // 1024}KB")
    print(f"블록 {len(blocks)}개, 그림 {images}장")


if __name__ == "__main__":
    main()
