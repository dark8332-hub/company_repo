"""반입 런북의 세 형식(md·docx·html)이 같은 내용인지.

마크다운이 원본이고 docx·html 은 tools/build_runbook.py 가 만든다. 산출물을 다시 만들지 않고
마크다운만 고치면 사이트에 건네지는 Word 문서가 옛 버전을 안내한다 — 1.1.4 시절 실제로
NOPASSWD 를 안내하는 .docx 가 남아 있었다. mtime 은 git 이 보존하지 않으므로 내용으로 대조한다:
제목 목록, 그림 수와 캡션, 코드 블록 수, 그리고 문서가 참조하는 그림 파일의 존재.

HTML 은 폐쇄망에서 열리므로 밖으로 나가는 참조가 없어야 한다.
"""
import re
import sys
from pathlib import Path

import docx
import pytest

PROJECT = Path(__file__).resolve().parent.parent
DOCS = PROJECT / "docs"
SOURCE = DOCS / "폐쇄망-반입-런북.md"
sys.path.insert(0, str(PROJECT / "tools"))
import build_runbook  # noqa: E402


@pytest.fixture(scope="module")
def blocks():
    return build_runbook.parse(SOURCE.read_text(encoding="utf-8"))


def headings(blocks, max_level=4):
    return [build_runbook.plain(text) for kind, (level, text) in
            ((k, b) for k, b in blocks if k == "heading") if level <= max_level]


def test_every_image_exists(blocks):
    missing = [src for kind, (alt, src) in ((k, b) for k, b in blocks if k == "image")
               if not (DOCS / src).exists()]
    assert not missing, f"런북이 참조하는 그림이 없습니다: {missing}"


def test_every_image_has_a_caption(blocks):
    bare = [src for kind, (alt, src) in ((k, b) for k, b in blocks if k == "image") if not alt.strip()]
    assert not bare, f"캡션 없는 그림: {bare}"


def test_docx_matches_markdown(blocks):
    document = docx.Document(str(SOURCE.with_suffix(".docx")))
    docx_headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
    assert docx_headings == headings(blocks), "docx 의 제목이 md 와 다릅니다. tools/build_runbook.py 를 다시 실행하세요."
    images = [alt for k, (alt, _) in ((k, b) for k, b in blocks if k == "image")]
    assert len(document.inline_shapes) == len(images)
    captions = [p.text for p in document.paragraphs if re.match(r"^그림 \d+\. ", p.text)]
    assert captions == [f"그림 {n}. {alt}" for n, alt in enumerate(images, 1)]
    assert len(document.tables) == sum(1 for k, _ in blocks if k == "table")


def test_html_matches_markdown(blocks):
    page = SOURCE.with_suffix(".html").read_text(encoding="utf-8")
    for title in headings(blocks, max_level=2):
        assert f">{build_runbook.inline_html(title)}</h" in page, f"html 에 제목이 없습니다: {title}"
    images = [alt for k, (alt, _) in ((k, b) for k, b in blocks if k == "image")]
    assert page.count("<figure>") == len(images)
    for n, alt in enumerate(images, 1):
        assert f"그림 {n}. {build_runbook.inline_html(alt)}" in page
    assert page.count("<pre>") == sum(1 for k, _ in blocks if k == "code")


def test_html_is_self_contained():
    page = SOURCE.with_suffix(".html").read_text(encoding="utf-8")
    external = re.findall(r'(?:src|href)="(https?://[^"]+)"', page)
    assert not external, f"폐쇄망 문서가 외부 자원을 참조합니다: {external}"
    assert "<script" not in page
    assert 'src="data:image/' in page, "그림이 data URI 로 들어 있어야 합니다"


def test_markdown_only_uses_supported_constructs():
    """생성기가 모르는 구문은 조용히 본문으로 흘러가므로 미리 잡는다."""
    text = SOURCE.read_text(encoding="utf-8")
    assert "<br" not in text and "<img" not in text and "<table" not in text
    assert not re.search(r"^\s*\d+\)\s", text, re.M), "번호 목록은 '1.' 형식만 지원합니다"
    assert not re.search(r"^#{5,}\s", text, re.M), "제목은 4단계까지만 지원합니다"
    assert not re.search(r"\[[^\]]+\]\((?!images/)[^)]+\)", text.replace("![", "")), "링크 구문은 지원하지 않습니다(그림만)"
