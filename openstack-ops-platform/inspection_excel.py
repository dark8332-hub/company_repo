"""일일점검 내용 Excel(xlsx) 생성기 (openpyxl).

`POST /api/reports/inspection.pdf` 와 동일한 payload 를 받아 운영자의 기존 점검표 양식
(번호 붙은 섹션 제목 + 점검 분류 | 점검 사항 | 점검 방법 | 점검 결과 | 특이사항 표)을
단일 시트 "일일점검" 으로 재현한다.
"""
from __future__ import annotations

import re
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

from inspection_report import STATUS_LABELS, _format_duration

SHEET_TITLE = "일일점검"
COLUMNS = ("점검 분류", "점검 사항", "점검 방법", "점검 결과", "특이사항", "직전 대비")
COL_WIDTHS = (14, 22, 46, 30, 40, 14)
MAX_CELL_CHARS = 32_000  # Excel 셀 한도 32,767

HEADER_FILL = PatternFill("solid", fgColor="2F5F8F")
WARNING_FILL = PatternFill("solid", fgColor="FFF4E3")
UNAVAILABLE_FILL = PatternFill("solid", fgColor="F1F4F8")
ROW_FILLS = {"warning": WARNING_FILL, "unavailable": UNAVAILABLE_FILL}
STATUS_FONT_COLORS = {"healthy": "1F8F54", "excepted": "1F8F54", "warning": "D67918", "unavailable": "6B7B93", "pending": "8996A8", "skipped": "8996A8"}
# 점검 결과 셀 앞에 붙는 상태 라벨. 요약 표기와 맞추기 위해 skipped/excepted 는 짧게 표기한다.
RESULT_STATUS_LABELS = {**STATUS_LABELS, "skipped": "제외", "excepted": "정상(예외 처리)"}

THIN = Side(style="thin", color="8A8A8A")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NAVY = "12264A"
TEXT = "333F52"
MUTED = "6B7B93"
FONT_NAME = "맑은 고딕"

TITLE_FONT = Font(name=FONT_NAME, size=16, bold=True, color=NAVY)
META_LABEL_FONT = Font(name=FONT_NAME, size=9, color=MUTED)
META_VALUE_FONT = Font(name=FONT_NAME, size=9, color=TEXT)
SECTION_FONT = Font(name=FONT_NAME, size=13, bold=True, color=NAVY)
HEADER_FONT = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name=FONT_NAME, size=9, color=TEXT)
BODY_BOLD_FONT = Font(name=FONT_NAME, size=9, bold=True, color="25334A")

WRAP_TOP = Alignment(wrap_text=True, vertical="top")
WRAP_CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")


def _clean(value, limit: int = MAX_CELL_CHARS) -> str:
    """openpyxl 이 거부하는 제어 문자를 제거하고 Excel 셀 한도 이내로 자른다."""
    text = str(value if value is not None else "").replace("\r\n", "\n").replace("\r", "\n")
    text = ILLEGAL_CHARACTERS_RE.sub("", text)
    if len(text) > limit:
        text = text[: limit - 1].rstrip() + "…"
    return text


def _safe_cell(value) -> str:
    """수식으로 해석되지 않도록 '=' 등으로 시작하는 문자열 앞에 공백을 붙인다."""
    text = _clean(value)
    return " " + text if text[:1] in ("=", "+", "-", "@") and len(text) > 1 else text


def _parse_checked_at(value) -> datetime | None:
    """app.js 가 보내는 ko-KR 형식("2026. 9. 3. 오후 2:05"), ISO, 'YYYY-MM-DD HH:MM' 을 해석한다."""
    text = str(value or "").strip()
    if not text:
        return None
    match = re.match(r"^(\d{4})\.\s*(\d{1,2})\.\s*(\d{1,2})\.\s*(오전|오후)?\s*(\d{1,2}):(\d{2})", text)
    if match:
        year, month, day, meridiem, hour, minute = match.groups()
        hour_value = int(hour)
        if meridiem == "오후" and hour_value < 12:
            hour_value += 12
        elif meridiem == "오전" and hour_value == 12:
            hour_value = 0
        try:
            return datetime(int(year), int(month), int(day), hour_value, int(minute))
        except ValueError:
            return None
    for candidate in (text, text.replace("Z", "+00:00")):
        try:
            return datetime.fromisoformat(candidate)
        except ValueError:
            pass
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass
    return None


def inspection_workbook_filename(payload: dict) -> str:
    """app.js 의 PDF 파일명 규칙(일일점검_<provider>_<YYYYMMDD-HHMM>.pdf)을 xlsx 로 그대로 옮긴 것."""
    provider = str(payload.get("provider") or "-")
    provider = re.sub(r'[\\/:*?"<>|\s]+', "_", provider) or "-"
    checked_at = _parse_checked_at(payload.get("checked_at")) or datetime.now()
    return f"일일점검_{provider}_{checked_at.strftime('%Y%m%d-%H%M')}.xlsx"


def _write_row(ws, row_index: int, values, font, fill=None, alignment=WRAP_TOP, fonts=None):
    for offset, value in enumerate(values, start=1):
        cell = ws.cell(row=row_index, column=offset, value=value)
        cell.font = (fonts or {}).get(offset, font)
        cell.alignment = alignment
        cell.border = BORDER
        if fill is not None:
            cell.fill = fill


def _estimate_row_height(values, widths=COL_WIDTHS, line_height: float = 13.0, minimum: float = 18.0, maximum: float = 409.0) -> float:
    """줄바꿈·열 너비를 기준으로 행 높이를 어림한다 (Excel 은 wrap 행 높이를 자동 계산하지 않는 뷰어가 많다)."""
    lines = 1
    for value, width in zip(values, widths):
        text = str(value or "")
        chars_per_line = max(4, int(width * 1.1))
        count = 0
        for line in text.split("\n"):
            count += max(1, -(-len(line) // chars_per_line))
        lines = max(lines, count)
    return min(maximum, max(minimum, lines * line_height + 5))


def build_inspection_workbook(payload: dict) -> bytes:
    provider = _clean(payload.get("provider") or "-", 200)
    checked_at = _clean(payload.get("checked_at") or "최근 결과", 100)
    trigger = "예약" if payload.get("trigger") == "scheduled" else "수동"
    duration = _format_duration(payload.get("duration_seconds")) if payload.get("duration_seconds") is not None else "-"
    generated_at = datetime.now().strftime("%Y-%m-%d %H:%M")
    groups = [group for group in payload.get("groups") or [] if group.get("rows")]
    rows_all = [row for group in groups for row in group["rows"]]

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_TITLE
    for index, width in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    last_col = len(COLUMNS)

    # 제목 블록
    ws.cell(row=1, column=1, value="OpenStack 일일점검").font = TITLE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.row_dimensions[1].height = 28
    meta = (("Provider", provider), ("점검 시각", checked_at), ("실행 구분", trigger), ("소요 시간", duration), ("생성 시각", generated_at))
    row_index = 2
    for label, value in meta:
        label_cell = ws.cell(row=row_index, column=1, value=label)
        label_cell.font = META_LABEL_FONT
        label_cell.alignment = Alignment(vertical="center")
        value_cell = ws.cell(row=row_index, column=2, value=_safe_cell(value))
        value_cell.font = META_VALUE_FONT
        value_cell.alignment = Alignment(vertical="center")
        ws.merge_cells(start_row=row_index, start_column=2, end_row=row_index, end_column=last_col)
        ws.row_dimensions[row_index].height = 16
        row_index += 1

    # 요약 행
    counts: dict[str, int] = {}
    for row in rows_all:
        counts[str(row.get("status") or "")] = counts.get(str(row.get("status") or ""), 0) + 1
    summary = (
        ("전체", len(rows_all), NAVY),
        ("정상", counts.get("healthy", 0) + counts.get("excepted", 0), STATUS_FONT_COLORS["healthy"]),
        ("주의", counts.get("warning", 0), STATUS_FONT_COLORS["warning"]),
        ("확인 불가", counts.get("unavailable", 0), STATUS_FONT_COLORS["unavailable"]),
        ("수집 대기", counts.get("pending", 0), STATUS_FONT_COLORS["pending"]),
        ("제외", counts.get("skipped", 0), STATUS_FONT_COLORS["skipped"]),
    )
    row_index += 1
    summary_label_row, summary_value_row = row_index, row_index + 1
    for column, (label, value, color) in enumerate(summary, start=1):
        label_cell = ws.cell(row=summary_label_row, column=column, value=label)
        label_cell.font = Font(name=FONT_NAME, size=9, bold=True, color="FFFFFF")
        label_cell.fill = HEADER_FILL
        label_cell.alignment = WRAP_CENTER
        label_cell.border = BORDER
        value_cell = ws.cell(row=summary_value_row, column=column, value=value)
        value_cell.font = Font(name=FONT_NAME, size=12, bold=True, color=color)
        value_cell.alignment = WRAP_CENTER
        value_cell.border = BORDER
    ws.row_dimensions[summary_label_row].height = 18
    ws.row_dimensions[summary_value_row].height = 22
    row_index = summary_value_row + 1
    freeze_row = row_index + 1  # 제목·요약 아래에서 틀 고정

    # 섹션별 표
    for number, group in enumerate(groups, start=1):
        rows = group["rows"]
        row_index += 1  # 빈 줄
        section_cell = ws.cell(row=row_index, column=1, value=f"{number}. {_clean(group.get('title') or '-', 200)}")
        section_cell.font = SECTION_FONT
        section_cell.alignment = Alignment(vertical="center")
        section_cell.border = Border(bottom=Side(style="medium", color="2F5F8F"))
        ws.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=last_col)
        for column in range(2, last_col + 1):
            ws.cell(row=row_index, column=column).border = Border(bottom=Side(style="medium", color="2F5F8F"))
        ws.row_dimensions[row_index].height = 24
        row_index += 1

        _write_row(ws, row_index, COLUMNS, HEADER_FONT, fill=HEADER_FILL, alignment=WRAP_CENTER)
        ws.row_dimensions[row_index].height = 20
        row_index += 1

        for row in rows:
            status = str(row.get("status") or "")
            label = RESULT_STATUS_LABELS.get(status, status or "-")
            result_text = _clean(row.get("result"))
            result_value = f"[{label}] {result_text}" if result_text else f"[{label}]"
            values = (
                _safe_cell(row.get("category")),
                _safe_cell(row.get("name")),
                _safe_cell(row.get("method")),
                _safe_cell(result_value),
                _safe_cell(row.get("note")),
                _safe_cell(row.get("change")),
            )
            fonts = {2: BODY_BOLD_FONT}
            status_color = STATUS_FONT_COLORS.get(status)
            if status_color:
                fonts[4] = Font(name=FONT_NAME, size=9, color=status_color, bold=status in ("warning", "unavailable"))
            _write_row(ws, row_index, values, BODY_FONT, fill=ROW_FILLS.get(status), fonts=fonts)
            ws.row_dimensions[row_index].height = _estimate_row_height(values)
            row_index += 1

    # 틀 고정·인쇄 설정
    ws.freeze_panes = ws.cell(row=freeze_row, column=1)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5)
    ws.print_options.horizontalCentered = True
    ws.oddFooter.left.text = f"OKESTRO OpenStack Operations Platform · {provider}"
    ws.oddFooter.right.text = "&P / &N"
    ws.sheet_view.showGridLines = False

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


if __name__ == "__main__":
    from pathlib import Path

    from openpyxl import load_workbook

    sample = {
        "provider": "hnti / prod",
        "checked_at": "2026. 9. 3. 오후 2:05",
        "trigger": "scheduled",
        "duration_seconds": 83.4,
        "groups": [
            {"title": "시스템 기본 점검", "rows": [
                {"category": "Resource", "name": "CPU 사용률", "method": "top", "status": "healthy", "result": "avg 12%", "note": "", "change": "변화 없음"},
                {"category": "Resource", "name": "Disk 사용률", "method": "df -h", "status": "warning", "result": "/var 91%", "note": "임계치 초과", "change": "악화"},
                {"category": "System", "name": "Mount 상태", "method": "ls -lh /var/lib/glance/images\nls -lh /var/lib/cinder/conversion", "status": "unavailable", "result": "", "note": "ssh timeout \x00\x01", "change": ""},
            ]},
            {"title": "Middleware 점검", "rows": [
                {"category": "Clustering", "name": "Pcs cluster", "method": "pcs status", "status": "pending", "result": "", "note": "", "change": ""},
                {"category": "Clustering", "name": "Mysql cluster", "method": "show status like 'wsrep_cluster_weight';", "status": "skipped", "result": "=SUM(1)", "note": "", "change": ""},
                {"category": "Clustering", "name": "Rabbitmq cluster", "method": "rabbitmqctl cluster_status", "status": "excepted", "result": "x" * 40_000, "note": "long", "change": "신규"},
            ]},
            {"title": "빈 그룹", "rows": []},
        ],
    }
    out = Path("/tmp/claude-0/-root/b2bfb18e-afce-497b-8247-91fb24a9fa0d/scratchpad/sample.xlsx")
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_bytes(build_inspection_workbook(sample))

    wb = load_workbook(out)
    assert wb.sheetnames == [SHEET_TITLE], wb.sheetnames
    ws = wb[SHEET_TITLE]
    assert ws["A1"].value == "OpenStack 일일점검"
    header_rows = [r for r in range(1, ws.max_row + 1) if tuple(ws.cell(row=r, column=c).value for c in range(1, 7)) == COLUMNS]
    assert len(header_rows) == 2, header_rows  # 빈 그룹은 제외
    section_titles = [ws.cell(row=r - 1, column=1).value for r in header_rows]
    assert section_titles == ["1. 시스템 기본 점검", "2. Middleware 점검"], section_titles
    assert ws.cell(row=header_rows[0], column=1).fill.fgColor.rgb.endswith("2F5F8F")
    assert ws.cell(row=header_rows[0], column=1).font.bold and ws.cell(row=header_rows[0], column=1).font.color.rgb.endswith("FFFFFF")
    # 표 본문 행 수: 첫 표 3행, 두 번째 표 3행
    body_1 = header_rows[1] - header_rows[0] - 3  # 헤더, 빈 줄, 섹션 제목 제외
    assert body_1 == 3, body_1
    assert ws.max_row - header_rows[1] == 3, ws.max_row - header_rows[1]
    assert ws.cell(row=header_rows[0] + 2, column=4).value.startswith("[주의] /var 91%")
    assert ws.cell(row=header_rows[0] + 2, column=1).fill.fgColor.rgb.endswith("FFF4E3")
    assert ws.cell(row=header_rows[0] + 3, column=1).fill.fgColor.rgb.endswith("F1F4F8")
    assert "\x00" not in ws.cell(row=header_rows[0] + 3, column=5).value
    assert len(ws.cell(row=header_rows[1] + 3, column=4).value) <= MAX_CELL_CHARS + 20
    assert ws.cell(row=header_rows[1] + 2, column=4).value.startswith("[제외] =SUM")
    summary_row = next(r for r in range(1, 12) if ws.cell(row=r, column=1).value == "전체")
    assert [ws.cell(row=summary_row + 1, column=c).value for c in range(1, 7)] == [6, 2, 1, 1, 1, 1]
    assert ws.freeze_panes == f"A{summary_row + 3}", ws.freeze_panes
    assert ws.page_setup.orientation == "landscape"
    assert inspection_workbook_filename(sample) == "일일점검_hnti_prod_20260903-1405.xlsx", inspection_workbook_filename(sample)
    print("OK", out, f"{out.stat().st_size} bytes, {ws.max_row} rows")
